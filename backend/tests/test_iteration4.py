"""
Iteration 4 backend tests:
- Contract Redline endpoints (redline, redline-from-review, redline-upload)
- Templates Library endpoints (list, get, generate, docx, share, public share view)
- Modules regression (templates + bulk-contracts present)
"""
import io
import os
import time
import uuid as _uuid
import requests
import pytest
from docx import Document as DocxDocument

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
if not BASE_URL:
    with open("/app/frontend/.env") as f:
        for line in f:
            if line.startswith("REACT_APP_BACKEND_URL"):
                BASE_URL = line.split("=", 1)[1].strip().rstrip("/")
API = f"{BASE_URL}/api"

TS = f"{int(time.time())}-{_uuid.uuid4().hex[:6]}"
EMAIL = f"iter4+{TS}@nyaya.in"
PW = "pass1234"

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@pytest.fixture(scope="module")
def auth_headers():
    r = requests.post(f"{API}/auth/register",
                      json={"name": "Iter4", "email": EMAIL, "password": PW},
                      timeout=30)
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['token']}"}


def _make_docx(text: str) -> bytes:
    doc = DocxDocument()
    for line in text.split("\n"):
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


SAMPLE_CONTRACT = """MUTUAL NON-DISCLOSURE AGREEMENT
This NDA is between Alpha Pvt Ltd (Party A) and Beta Pvt Ltd (Party B).
Both parties may share confidential information for evaluating a joint venture.
Term: 3 years. Governing law: India. Jurisdiction: Bengaluru.
Liability capped at INR 10 lakh. Standard carve-outs apply.
"""


# ----------------- Templates: list & get -----------------
class TestTemplatesList:
    def test_list_templates(self):
        r = requests.get(f"{API}/templates", timeout=20)
        assert r.status_code == 200
        data = r.json()
        assert "categories" in data and "templates" in data
        slugs = {t["slug"] for t in data["templates"]}
        expected = {
            "nda-mutual", "nda-oneway", "employment", "consultant",
            "rental", "commercial-lease", "freelance-sow", "mou",
            "founder", "power-of-attorney", "will", "assignment-ip",
        }
        assert expected.issubset(slugs), f"Missing: {expected - slugs}"
        assert len(data["templates"]) == 12
        cats = set(data["categories"])
        assert {"Business", "HR", "Personal", "Property"}.issubset(cats)

    def test_get_template_spec(self):
        r = requests.get(f"{API}/templates/nda-mutual", timeout=20)
        assert r.status_code == 200
        d = r.json()
        assert d["slug"] == "nda-mutual"
        assert d["title"].startswith("Mutual Non-Disclosure")
        keys = {f["key"] for f in d["fields"]}
        assert {"party_a_name", "party_b_name", "purpose",
                "duration_years", "jurisdiction_city"}.issubset(keys)
        # defaults preserved
        for f in d["fields"]:
            if f["key"] == "duration_years":
                assert f.get("default") == "3"
            if f["key"] == "jurisdiction_city":
                assert f.get("default") == "Bengaluru"

    def test_get_template_unknown(self):
        r = requests.get(f"{API}/templates/no-such-slug", timeout=20)
        assert r.status_code == 404


# ----------------- Templates: generate & docx & share -----------------
class TestTemplateFlows:
    def test_generate_requires_auth(self):
        r = requests.post(f"{API}/templates/nda-mutual/generate",
                          json={"inputs": {}}, timeout=30)
        assert r.status_code in (401, 403)

    def test_generate_missing_required_400(self, auth_headers):
        r = requests.post(f"{API}/templates/nda-mutual/generate",
                          headers=auth_headers, json={"inputs": {}}, timeout=30)
        assert r.status_code == 400
        assert "required" in r.text.lower()

    def test_generate_success(self, auth_headers):
        payload = {"inputs": {
            "party_a_name": "Alpha Pvt Ltd",
            "party_a_address": "123 MG Road, Bengaluru",
            "party_b_name": "Beta Pvt Ltd",
            "party_b_address": "456 Brigade Road, Bengaluru",
            "purpose": "Evaluate a fintech JV",
            "duration_years": "3",
            "jurisdiction_city": "Bengaluru",
        }}
        r = requests.post(f"{API}/templates/nda-mutual/generate",
                          headers=auth_headers, json=payload, timeout=120)
        assert r.status_code == 200, r.text
        d = r.json()
        assert d["slug"] == "nda-mutual"
        assert d["title"]
        assert isinstance(d["text"], str) and len(d["text"]) > 500
        assert "id" in d
        # store for share test
        pytest.template_text = d["text"]
        pytest.template_title = d["title"]

    def test_docx_download(self, auth_headers):
        payload = {"inputs": {
            "party_a_name": "Alpha Pvt Ltd",
            "party_a_address": "123 MG Road, Bengaluru",
            "party_b_name": "Beta Pvt Ltd",
            "party_b_address": "456 Brigade Road, Bengaluru",
            "purpose": "Evaluate a fintech JV",
            "duration_years": "3",
            "jurisdiction_city": "Bengaluru",
        }}
        r = requests.post(f"{API}/templates/nda-mutual/docx",
                          headers=auth_headers, json=payload, timeout=180)
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith(DOCX_MIME)
        assert "attachment" in r.headers.get("content-disposition", "").lower()
        content = r.content
        assert len(content) > 5000, f"docx too small: {len(content)}"
        # Parse with python-docx
        doc = DocxDocument(io.BytesIO(content))
        paras = [p.text for p in doc.paragraphs if p.text.strip()]
        assert len(paras) >= 3
        joined = "\n".join(paras)
        assert "Nyaya Mitra" in joined or "Non-Disclosure" in joined or "NDA" in joined.upper()

    def test_share_requires_auth(self):
        r = requests.post(f"{API}/templates/share",
                          json={"slug": "nda-mutual", "title": "T", "text": "x"},
                          timeout=20)
        assert r.status_code in (401, 403)

    def test_share_and_public_view(self, auth_headers):
        text = getattr(pytest, "template_text", "Sample draft body text here.")
        title = getattr(pytest, "template_title", "Mutual NDA")
        r = requests.post(f"{API}/templates/share", headers=auth_headers,
                          json={"slug": "nda-mutual", "title": title, "text": text},
                          timeout=20)
        assert r.status_code == 200, r.text
        share_id = r.json()["share_id"]
        assert isinstance(share_id, str) and len(share_id) == 12

        # Public read - no auth
        r2 = requests.get(f"{API}/templates/share/{share_id}", timeout=20)
        assert r2.status_code == 200
        d = r2.json()
        assert d["share_id"] == share_id
        assert d["slug"] == "nda-mutual"
        assert d["title"] == title
        assert d["text"] == text
        assert "user_name" in d
        assert "created_at" in d
        assert "user_id" not in d
        assert "_id" not in d

    def test_share_unknown_404(self):
        r = requests.get(f"{API}/templates/share/deadbeefbaad", timeout=20)
        assert r.status_code == 404


# ----------------- Redline endpoints -----------------
class TestRedline:
    def test_redline_no_auth(self):
        r = requests.post(f"{API}/review/contract/redline",
                          json={"contract_text": SAMPLE_CONTRACT}, timeout=30)
        assert r.status_code in (401, 403)

    def test_redline_from_review_no_auth(self):
        r = requests.post(f"{API}/review/contract/redline-from-review",
                          json={"contract_text": SAMPLE_CONTRACT, "review": {}}, timeout=30)
        assert r.status_code in (401, 403)

    def test_redline_upload_no_auth(self):
        files = {"file": ("c.docx", _make_docx(SAMPLE_CONTRACT), DOCX_MIME)}
        r = requests.post(f"{API}/review/contract/redline-upload",
                          files=files, timeout=30)
        assert r.status_code in (401, 403)

    def _check_redline_docx(self, content: bytes):
        assert len(content) > 5000
        doc = DocxDocument(io.BytesIO(content))
        paras = "\n".join(p.text for p in doc.paragraphs)
        assert "Nyaya Mitra" in paras
        assert "Contract Redline" in paras or "Redline Report" in paras
        assert "Executive Summary" in paras
        assert "Contract Text" in paras
        assert "Redline Annotations" in paras

    def test_redline_from_review(self, auth_headers):
        review = {
            "contract_type": "NDA",
            "overall_risk": "medium",
            "risk_score": 55,
            "summary": "Balanced NDA with standard clauses.",
            "red_flags": [
                {"clause": "Indemnification", "severity": "high",
                 "issue": "Uncapped indemnity", "suggestion": "Cap at fees paid"},
            ],
            "missing_clauses": [
                {"clause": "Force Majeure", "importance": "medium",
                 "suggestion": "Add force majeure clause"},
            ],
            "favourable_clauses": [
                {"clause": "Term", "note": "Reasonable 3-year term"}
            ],
        }
        r = requests.post(f"{API}/review/contract/redline-from-review",
                          headers=auth_headers,
                          json={"contract_text": SAMPLE_CONTRACT,
                                "review": review, "filename": "sample.docx"},
                          timeout=60)
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith(DOCX_MIME)
        cd = r.headers.get("content-disposition", "").lower()
        assert "attachment" in cd and ".docx" in cd
        self._check_redline_docx(r.content)

    def test_redline_from_text(self, auth_headers):
        r = requests.post(f"{API}/review/contract/redline",
                          headers=auth_headers,
                          json={"contract_text": SAMPLE_CONTRACT},
                          timeout=180)
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith(DOCX_MIME)
        self._check_redline_docx(r.content)

    def test_redline_upload(self, auth_headers):
        files = {"file": ("c.docx", _make_docx(SAMPLE_CONTRACT), DOCX_MIME)}
        r = requests.post(f"{API}/review/contract/redline-upload",
                          headers=auth_headers, files=files, timeout=180)
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith(DOCX_MIME)
        self._check_redline_docx(r.content)


# ----------------- Modules regression -----------------
class TestModules:
    def test_modules_has_templates_and_bulk(self, auth_headers):
        r = requests.get(f"{API}/modules", headers=auth_headers, timeout=20)
        assert r.status_code == 200
        s = str(r.json())
        assert "templates" in s
        assert "book-marked" in s
        assert "bulk-contracts" in s
