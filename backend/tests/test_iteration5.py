"""Iteration 5 tests — Regional-language translation for Contract Templates."""
import io
import os
import re
import time

import pytest
import requests

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
if not BASE_URL:
    # Read from frontend/.env
    try:
        with open("/app/frontend/.env") as f:
            for line in f:
                if line.startswith("REACT_APP_BACKEND_URL="):
                    BASE_URL = line.split("=", 1)[1].strip().rstrip("/")
                    break
    except Exception:
        pass
assert BASE_URL, "REACT_APP_BACKEND_URL missing"
API = f"{BASE_URL}/api"

TS = int(time.time())
EMAIL = f"trans{TS}@nyaya.in"
PASS = "pass1234"


@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{API}/auth/register", json={"name": "Trans Tester", "email": EMAIL, "password": PASS}, timeout=30)
    assert r.status_code in (200, 201), r.text
    tok = r.json().get("token") or r.json().get("access_token")
    if not tok:
        r2 = requests.post(f"{API}/auth/login", json={"email": EMAIL, "password": PASS}, timeout=30)
        tok = r2.json().get("token") or r2.json().get("access_token")
    assert tok, "no token"
    return tok


@pytest.fixture
def auth(token):
    return {"Authorization": f"Bearer {token}"}


SAMPLE_EN = (
    "MUTUAL NON-DISCLOSURE AGREEMENT\n\n"
    "This Agreement is entered into between ABC Tech Pvt Ltd and XYZ Analytics LLP on 12 January 2026. "
    "The parties agree to pay ₹35000 in consideration. Any disputes shall be governed by Section 138 of the "
    "Negotiable Instruments Act, 1881 and the Companies Act, 2013.\n\n"
    "1. Confidential Information means any non-public information exchanged.\n"
    "2. Term: 24 months from the effective date.\n"
)


# --- GET /api/languages -----------------------------------------------------
def test_languages_list():
    r = requests.get(f"{API}/languages", timeout=15)
    assert r.status_code == 200, r.text
    langs = r.json()
    codes = {l["code"] for l in langs}
    assert codes == {"english", "hindi", "tamil", "telugu", "bengali", "marathi"}, codes
    for l in langs:
        assert {"code", "label", "native", "script"}.issubset(l.keys())


# --- Auth requirement -------------------------------------------------------
def test_translate_requires_auth():
    r = requests.post(f"{API}/templates/translate",
                      json={"text": SAMPLE_EN, "target_language": "hindi"}, timeout=15)
    assert r.status_code in (401, 403), r.status_code


# --- Validation -------------------------------------------------------------
def test_translate_unsupported_language(auth):
    r = requests.post(f"{API}/templates/translate",
                      json={"text": SAMPLE_EN, "target_language": "french"}, headers=auth, timeout=15)
    assert r.status_code == 400
    assert "unsupported" in r.text.lower()


def test_translate_text_too_short(auth):
    r = requests.post(f"{API}/templates/translate",
                      json={"text": "hi", "target_language": "hindi"}, headers=auth, timeout=15)
    assert r.status_code == 400
    assert "short" in r.text.lower()


# --- English echo (no LLM call) --------------------------------------------
def test_translate_english_echo(auth):
    r = requests.post(f"{API}/templates/translate",
                      json={"text": SAMPLE_EN, "target_language": "english"}, headers=auth, timeout=15)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["language"] == "english"
    assert data["text"] == SAMPLE_EN


# --- 5 non-English translations --------------------------------------------
SCRIPT_RANGES = {
    "hindi":   (0x0900, 0x097F),   # Devanagari
    "marathi": (0x0900, 0x097F),   # Devanagari
    "tamil":   (0x0B80, 0x0BFF),
    "telugu":  (0x0C00, 0x0C7F),
    "bengali": (0x0980, 0x09FF),
}
NATIVE = {"hindi": "हिंदी", "tamil": "தமிழ்", "telugu": "తెలుగు", "bengali": "বাংলা", "marathi": "मराठी"}


def _has_script(text: str, lo: int, hi: int) -> bool:
    return any(lo <= ord(ch) <= hi for ch in text)


@pytest.mark.parametrize("lang", ["hindi", "tamil", "telugu", "bengali", "marathi"])
def test_translate_indic_language(auth, lang):
    r = requests.post(f"{API}/templates/translate",
                      json={"text": SAMPLE_EN, "target_language": lang},
                      headers=auth, timeout=90)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["language"] == lang
    assert data["native"] == NATIVE[lang]
    text = data["text"]
    assert text and len(text.strip()) > 20, f"empty translation for {lang}: {text!r}"
    lo, hi = SCRIPT_RANGES[lang]
    assert _has_script(text, lo, hi), f"{lang} response has no target script chars: {text[:200]!r}"
    # Proper nouns preserved verbatim
    for token in ["ABC Tech Pvt Ltd", "XYZ Analytics LLP", "35000", "138"]:
        assert token in text, f"{lang} translation dropped `{token}`: {text[:400]!r}"


# --- DOCX-from-text ---------------------------------------------------------
def test_docx_from_text_hindi(auth):
    # Get a small Hindi translation first
    tr = requests.post(f"{API}/templates/translate",
                       json={"text": SAMPLE_EN, "target_language": "hindi"},
                       headers=auth, timeout=90).json()
    hi_text = tr["text"]
    r = requests.post(f"{API}/templates/docx-from-text",
                      json={"title": "Mutual NDA", "text": hi_text, "language": "hindi"},
                      headers=auth, timeout=30)
    assert r.status_code == 200, r.text
    assert r.headers.get("content-type", "").startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    body = r.content
    assert len(body) >= 5000, f"docx too small: {len(body)} bytes"
    cd = r.headers.get("content-disposition", "")
    assert "Hindi" in cd, cd
    # Open via python-docx
    from docx import Document
    doc = Document(io.BytesIO(body))
    all_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Mutual NDA" in all_text
    assert "Hindi" in all_text  # title suffixed with — Hindi
    # Devanagari chars present
    assert _has_script(all_text, 0x0900, 0x097F), "no Devanagari in DOCX"


def test_docx_from_text_english_no_suffix(auth):
    r = requests.post(f"{API}/templates/docx-from-text",
                      json={"title": "Sample", "text": SAMPLE_EN, "language": "english"},
                      headers=auth, timeout=30)
    assert r.status_code == 200
    from docx import Document
    doc = Document(io.BytesIO(r.content))
    txt = "\n".join(p.text for p in doc.paragraphs)
    # English title should NOT have "— English" suffix per server logic
    assert "Sample" in txt
    assert "— English" not in txt


def test_docx_from_text_missing_text(auth):
    r = requests.post(f"{API}/templates/docx-from-text",
                      json={"title": "x", "text": "   ", "language": "hindi"},
                      headers=auth, timeout=15)
    assert r.status_code == 400


